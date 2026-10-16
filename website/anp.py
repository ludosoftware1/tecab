"""Histórico de movimentação do terminal (Resolução ANP nº 881/2022, art. 26, III, d).

Leitura e geração da planilha no layout da ANP (abas "Metadados" e "Dados", nomes dos campos na linha 3 e
dados a partir da linha 4) e dados agregados para o painel de gestão.
"""
import datetime
import io
import re
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.db import transaction
from django.utils import timezone

from .models import MovimentacaoANP, ProdutoANP

# Terminal TECAB na tabela T008 do SIMP. Os campos são iguais em todas as linhas da planilha.
TERMINAL = {'codigo_anp': 1033585, 'nome': 'TERMINAL TECAB', 'municipio': 'CABEDELO', 'uf': 'PB'}

CAMPOS = ['mes_de_referencia', 'codigo_anp_do_terminal', 'nome_do_terminal', 'municipio_do_terminal', 'uf',
          'sentido_da_operacao', 'tipo_da_operacao', 'modo_de_transporte', 'codigo_anp_do_produto',
          'descricao_do_produto', 'volume_m3']
LINHA_CAMPOS = 3
MODELO_XLSX = Path(__file__).resolve().parent / 'data' / 'modelo-movimentacao-anp.xlsx'

# A publicação do mês anterior deve estar no site até o dia 15 de cada mês.
DIA_LIMITE_PUBLICACAO = 15

_MES_RE = re.compile(r'^(\d{4})-(\d{2})$')


class PlanilhaInvalida(Exception):
    pass


def mes_para_data(valor):
    """'2026-09' (ou uma data) -> date(2026, 9, 1)."""
    if isinstance(valor, datetime.datetime):
        valor = valor.date()
    if isinstance(valor, datetime.date):
        return valor.replace(day=1)
    m = _MES_RE.match(str(valor or '').strip())
    if not m or not 1 <= int(m.group(2)) <= 12:
        raise ValueError('mês de referência inválido: %r (use AAAA-MM)' % valor)
    return datetime.date(int(m.group(1)), int(m.group(2)), 1)


def mes_seguinte(data, n=1):
    total = data.year * 12 + data.month - 1 + n
    return datetime.date(total // 12, total % 12 + 1, 1)


def volume_decimal(valor):
    if isinstance(valor, float):
        valor = repr(valor)
    try:
        volume = Decimal(str(valor).strip().replace(',', '.')).quantize(Decimal('0.001'))
    except (InvalidOperation, ValueError):
        raise ValueError('volume inválido: %r' % valor)
    if volume < 0:
        raise ValueError('volume negativo: %s' % volume)
    return volume


# --------------------------------------------------------------------------- leitura
def ler_linhas(registros):
    """Valida as linhas da planilha (dicionários com os CAMPOS) e devolve (linhas, erros)."""
    linhas, erros = [], []
    validos = {
        'sentido': dict(MovimentacaoANP.SENTIDOS), 'tipo': dict(MovimentacaoANP.TIPOS),
        'modo': dict(MovimentacaoANP.MODOS),
    }
    for numero, reg in registros:
        try:
            mes = mes_para_data(reg['mes_de_referencia'])
            sentido = int(reg['sentido_da_operacao'])
            tipo = int(reg['tipo_da_operacao'])
            modo = int(reg['modo_de_transporte'])
            for nome, valor in (('sentido', sentido), ('tipo', tipo), ('modo', modo)):
                if valor not in validos[nome]:
                    raise ValueError('%s da operação inválido: %s' % (nome, valor))
            codigo = int(reg['codigo_anp_do_produto'] or 0)
            descricao = str(reg['descricao_do_produto'] or '').strip()
            if not descricao:
                raise ValueError('descrição do produto vazia')
            terminal = reg.get('codigo_anp_do_terminal')
            if terminal not in (None, '') and int(terminal) != TERMINAL['codigo_anp']:
                raise ValueError('código do terminal %s diferente do TECAB (%s)' % (terminal, TERMINAL['codigo_anp']))
            linhas.append({'mes': mes, 'sentido': sentido, 'tipo': tipo, 'modo': modo, 'codigo': codigo,
                           'descricao': descricao, 'volume': volume_decimal(reg['volume_m3'])})
        except (KeyError, TypeError, ValueError) as exc:
            erros.append('Linha %s: %s' % (numero, exc))
    return linhas, erros


def ler_xlsx(arquivo):
    """Lê a aba "Dados" de uma planilha no layout da ANP. Aceita caminho ou arquivo enviado."""
    import openpyxl  # importado aqui: só é necessário na importação/exportação

    try:
        wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    except Exception as exc:
        raise PlanilhaInvalida('Não foi possível abrir o arquivo como planilha Excel (.xlsx): %s' % exc)
    try:
        if 'Dados' not in wb.sheetnames:
            raise PlanilhaInvalida('A planilha não tem a aba "Dados" do layout da ANP.')
        ws = wb['Dados']
        cabecalho = None
        registros = []
        for numero, linha in enumerate(ws.iter_rows(values_only=True), start=1):
            if numero < LINHA_CAMPOS:
                continue
            if numero == LINHA_CAMPOS:
                cabecalho = [str(v).strip() if v is not None else '' for v in linha[:len(CAMPOS)]]
                if cabecalho != CAMPOS:
                    raise PlanilhaInvalida('Os nomes dos campos na linha 3 da aba "Dados" não seguem o layout da ANP '
                                           '(esperado: %s).' % ', '.join(CAMPOS))
                continue
            valores = list(linha[:len(CAMPOS)])
            if all(v is None or str(v).strip() == '' for v in valores):
                continue
            valores += [None] * (len(CAMPOS) - len(valores))
            registros.append((numero, dict(zip(CAMPOS, valores))))
        if cabecalho is None:
            raise PlanilhaInvalida('A aba "Dados" está vazia.')
    finally:
        wb.close()
    return ler_linhas(registros)


@transaction.atomic
def importar(linhas, usuario=None):
    """Grava as linhas substituindo, mês a mês, o que já existe nos meses presentes na planilha."""
    meses = sorted({l['mes'] for l in linhas})
    removidas, _ = MovimentacaoANP.objects.filter(mes_referencia__in=meses).delete()
    produtos = {}
    for l in linhas:
        chave = (l['codigo'], l['descricao'])
        if chave not in produtos:
            produtos[chave], _ = ProdutoANP.objects.get_or_create(codigo_anp=l['codigo'], descricao=l['descricao'])
    MovimentacaoANP.objects.bulk_create([
        MovimentacaoANP(mes_referencia=l['mes'], sentido=l['sentido'], tipo_operacao=l['tipo'],
                        modo_transporte=l['modo'], produto=produtos[(l['codigo'], l['descricao'])],
                        volume_m3=l['volume'], criado_por=usuario, atualizado_por=usuario)
        for l in sorted(linhas, key=lambda l: l['mes'])  # sort estável: mantém a ordem dentro do mês
    ])
    return {'linhas': len(linhas), 'meses': len(meses), 'removidas': removidas,
            'primeiro': meses[0] if meses else None, 'ultimo': meses[-1] if meses else None}


# --------------------------------------------------------------------------- exportação
def data_atualizacao():
    ultima = MovimentacaoANP.objects.order_by('-atualizado_em').values_list('atualizado_em', flat=True).first()
    return timezone.localtime(ultima).date() if ultima else None


def exportar_xlsx():
    """Planilha no layout da ANP, pronta para publicação no site."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Side

    wb = openpyxl.load_workbook(MODELO_XLSX)
    ws = wb['Dados']
    atualizado = data_atualizacao()
    ws['B1'] = atualizado.strftime('%d/%m/%Y') if atualizado else ''
    fino = Side(style='thin')
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    centro = Alignment(horizontal='center')
    linha = LINHA_CAMPOS + 1
    for m in MovimentacaoANP.objects.select_related('produto'):
        valores = [m.mes, TERMINAL['codigo_anp'], TERMINAL['nome'], TERMINAL['municipio'], TERMINAL['uf'],
                   m.sentido, m.tipo_operacao, m.modo_transporte, m.produto.codigo_anp, m.produto.descricao,
                   float(m.volume_m3)]
        for coluna, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha, column=coluna, value=valor)
            if coluna > 1:
                celula.border = borda
            if coluna in (2, 6, 7, 8, 9, 11):
                celula.alignment = centro
        ws.cell(row=linha, column=2).number_format = '0'
        ws.cell(row=linha, column=9).number_format = '000000000'
        ws.cell(row=linha, column=11).number_format = '0.00'
        linha += 1
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def linhas_csv():
    yield CAMPOS
    for m in MovimentacaoANP.objects.select_related('produto'):
        yield [m.mes, TERMINAL['codigo_anp'], TERMINAL['nome'], TERMINAL['municipio'], TERMINAL['uf'], m.sentido,
               m.tipo_operacao, m.modo_transporte, '%09d' % m.produto.codigo_anp, m.produto.descricao,
               ('%.3f' % m.volume_m3).replace('.', ',')]


# --------------------------------------------------------------------------- painel
def situacao_publicacao(hoje=None):
    """Prazo de publicação: o mês anterior deve estar lançado até o dia 15."""
    hoje = hoje or timezone.localdate()
    mes_atual = hoje.replace(day=1)
    devido = mes_seguinte(mes_atual, -1)                     # mês que já deveria estar (ou estar sendo) publicado
    prazo = mes_atual.replace(day=DIA_LIMITE_PUBLICACAO)
    lancado = MovimentacaoANP.objects.filter(mes_referencia=devido).exists()
    ultimo = MovimentacaoANP.objects.order_by('-mes_referencia').values_list('mes_referencia', flat=True).first()
    if lancado:
        proximo = mes_atual
        return {'status': 'em_dia', 'mes': devido, 'proximo_mes': proximo,
                'proximo_prazo': mes_seguinte(mes_atual).replace(day=DIA_LIMITE_PUBLICACAO), 'ultimo': ultimo}
    status = 'atrasado' if hoje > prazo else 'pendente'
    return {'status': status, 'mes': devido, 'prazo': prazo, 'dias': (prazo - hoje).days, 'ultimo': ultimo}


def pontos_de_atencao(movimentacoes):
    """Inconsistências que merecem revisão antes da publicação."""
    avisos = []
    chaves = Counter((m.mes_referencia, m.sentido, m.tipo_operacao, m.modo_transporte, m.produto_id)
                     for m in movimentacoes)
    nomes = {}
    for m in movimentacoes:
        nomes[m.produto_id] = m.produto.descricao
    for (mes, sentido, tipo, modo, produto), n in sorted(chaves.items()):
        if n > 1:
            avisos.append({
                'mes': mes, 'tipo': 'duplicado',
                'texto': '%d lançamentos com a mesma combinação: %s · %s · %s · %s.' % (
                    n, dict(MovimentacaoANP.SENTIDOS)[sentido], dict(MovimentacaoANP.TIPOS)[tipo],
                    dict(MovimentacaoANP.MODOS)[modo], nomes[produto]),
            })
    meses = sorted({m.mes_referencia for m in movimentacoes})
    if meses:
        esperado = meses[0]
        presentes = set(meses)
        while esperado <= meses[-1]:
            if esperado not in presentes:
                avisos.append({'mes': esperado, 'tipo': 'lacuna', 'texto': 'Mês sem nenhum lançamento.'})
            esperado = mes_seguinte(esperado)
    por_mes = defaultdict(lambda: defaultdict(Decimal))
    for m in movimentacoes:
        por_mes[m.mes_referencia][m.sentido] += m.volume_m3
    for mes, sentidos in sorted(por_mes.items()):
        for sentido, nome in MovimentacaoANP.SENTIDOS:
            if not sentidos.get(sentido):
                avisos.append({'mes': mes, 'tipo': 'sem_sentido', 'texto': 'Nenhum volume de %s no mês.' % nome.lower()})
    return sorted(avisos, key=lambda a: a['mes'], reverse=True)


def dados_painel(movimentacoes, produtos):
    """Estrutura compacta enviada ao navegador (o painel agrega e filtra no cliente)."""
    usados = {m.produto_id for m in movimentacoes}
    return {
        'produtos': [{'id': p.pk, 'nome': p.descricao, 'codigo': '%09d' % p.codigo_anp}
                     for p in produtos if p.pk in usados or p.ativo],
        'sentidos': dict(MovimentacaoANP.SENTIDOS),
        'tipos': dict(MovimentacaoANP.TIPOS),
        'modos': dict(MovimentacaoANP.MODOS),
        # [id, mês, sentido, tipo, modo, produto, volume]
        'linhas': [[m.pk, m.mes, m.sentido, m.tipo_operacao, m.modo_transporte, m.produto_id, float(m.volume_m3)]
                   for m in movimentacoes],
    }
